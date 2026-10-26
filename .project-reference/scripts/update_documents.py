from pathlib import Path
import json,shutil
from docx import Document
from docx.shared import Pt,RGBColor
from docx.enum.style import WD_STYLE_TYPE
ROOT=Path(__file__).resolve().parents[2]
backup=ROOT/'.project-reference/original-documents-2026-10-22'
backup.mkdir(exist_ok=True)
patches={
'DiscoverSalah_Proposal.docx':{
68:'Current local release: a responsive Urdu/English React web application with participating mosque lists, location controls, Nearby and Whole City discovery, daily and dated Jama’at schedules, public reporting and external Google Maps directions. No embedded live map, route duration, arrival margin or arrive-before-Jama’at ranking is enabled. Some newly integrated account/tracking text still needs Urdu review.',
69:'Management uses scoped representative, focal-person and administrator workspaces backed by the API and MongoDB. Registration, private Imam contacts, immutable timing submissions, independent review, in-app notifications, reassignment, escalation and audit history are implemented. Future-departure forecasts, native apps, paid routing and external SMS/email delivery are outside the current local release.',
71:'Implemented architecture: React and TypeScript frontend; Node.js and Express modular API; persistent transaction-capable MongoDB; separate outbox and retention worker. Modules separate authentication, venues, timetables, community, notices and administration. Server commands enforce authorization and persistence; browser fixtures are not runtime records.',
72:'Selected navigation approach: Google Maps URLs open externally after a user click, without an API key, billing account, embedded SDK or route API. The URL contains the published entrance coordinates and walking or two-wheeler mode, never the visitor’s precise origin. Google Maps handles the starting point and supported directions. DiscoverSalah receives no route duration or navigation result.',
74:'Approval refers to the reviewed immutable version. Transactions, current-account checks, expected revisions, idempotency and durable outbox records protect concurrent writes. The current compatibility workspace is bounded to 200 venues and 1,000 proposals with smaller supporting feeds; per-screen pagination is required before larger deployment. Route caching and clustered embedded maps remain deferred design.',
81:'Accounts use contact-number/password authentication. Only Admin issues focal-person credentials; only the responsible focal person issues representative credentials in their assigned areas. The issuer independently confirms identity and contact details. Temporary passwords expire after 24 hours and require a first-use change. Issuer resets and contact-number changes revoke sessions. No SMS OTP or verified email is required. Administrator TOTP and single-use recovery codes provide MFA.',
82:'Sensitive contacts and evidence are encrypted and excluded from public DTOs. Precise visitor location remains in memory and is not included in external Maps links. Origin/CSRF checks, body limits, rate limits and server schemas guard commands. Closed rejected/withdrawn applications are redacted after 90 days; production backups, restore rehearsal and key rotation remain outstanding.'},
'DiscoverSalah_Phase_3_DTS.docx':{
78:'The owner’s 22 October 2026 decision replaces the earlier Google Maps/Routes API selection with external Google Maps URLs for the current local release. No route API, embedded map, geocoding service, API key or Google billing account is needed. Current credentials and deployment follow the revision page at the start of this document. Earlier route calculations remain deferred design.',
107:'Implemented stack: React and TypeScript, Vite, Node.js/Express TypeScript API, native MongoDB driver and a separate worker. The configured persistent MongoDB replica set is connected. Dependencies are recorded in repository lockfiles. The original standalone HTML remains a visual reference.',
108:'Current deployment is local: frontend http://127.0.0.1:5173, API on loopback port 4000 and a separately running notification/retention worker. The frontend proxies /api to the backend. MongoDB uses the private backend environment configuration. No Vercel or other production hosting has been provisioned. Production HTTPS, supervision, backup automation and restore acceptance remain future work.',
111:'Use no-key Google Maps URLs for external directions. The URL contains a validated published entrance and travelmode=walking or travelmode=two-wheeler. Omit the visitor origin and private contacts. Open only after a user click with referrer/opener protection. Do not call a Google route API, infer travel duration or substitute a car route for motorcycle travel.',
131:'Current public behavior: show timetable and timing-confidence information, then offer external Google Maps directions. Directions remain available even when no session exists for the selected prayer/date. No Estimated reachable or Time is tight claim is calculated in this release. A future route adapter requires separate approval, implementation, cost controls and field validation.',
175:'Actor: management account owner. Sign in with the issuer-confirmed contact number and password. Administrators complete TOTP or single-use recovery-code MFA. A temporary password must be changed before management commands become available.',
229:'Actor: Admin. Review uncovered areas and assignments, issue focal credentials, manage authorized credential resets/contact changes and maintenance responsibility. Representative credentials are issued only by the responsible focal person, not directly by Admin.',
250:'Management accounts are individually issued. Only Admin creates focal credentials; only the responsible focal person creates representative credentials within assigned areas. Issuers confirm identity and contact details independently, then privately hand over a generated temporary password valid for 24 hours. The account owner must change it before managing records. No verified-email or SMS OTP onboarding/recovery channel is required.',
251:'Recovery follows the same issuer hierarchy. The issuer rechecks identity/contact before resetting credentials or changing the account contact number. These commands revoke sessions; temporary-password reset requires a new private password. Administrator sign-in uses locally verified TOTP with single-use offline recovery codes. Lost administrator MFA requires controlled operator recovery, never a public bypass.',
541:'Durable in-app notifications and review queues are implemented through a transactional outbox worker. External email/SMS/push delivery and reminder services are not enabled. Issued credentials are handed over privately; SMS/email invitation and recovery endpoints are retired.'},
'DiscoverSalah_SDS_Phase_2.docx':{
65:'The implemented local solution is a modular React/TypeScript web application, Node.js/Express API, background worker and persistent MongoDB replica set. Immutable submissions and independent reviews protect timing integrity. External Google Maps URLs provide navigation; embedded maps, routing APIs, arrival calculations and forecasts remain deferred.',
138:'Only Admin issues focal-person credentials. Only the responsible focal person issues local representative credentials in their assigned areas; Admin does not create representatives directly. Application approval alone grants no access. The issuer checks identity, contact, authority and scope, then privately provides a generated 24-hour temporary password. The recipient must change it on first login. No OTP invitation acceptance is used.',
209:'Account credentials and recovery',
210:'Opaque AuthSession records and HttpOnly cookies authenticate the account with authVersion checks, CSRF and origin validation. Administrator sign-in first creates a short-lived MFA challenge, then verifies TOTP or a single-use recovery code. Temporary credentials carry mustChangePassword and temporaryPasswordExpiresAt. Protected credential issuance/reset/contact-change commands enforce the issuer hierarchy and revoke sessions when needed. Legacy SMS invitation/recovery endpoints return 410.',
240:'Community and management screens use explicit versioned API commands; consult backend route schemas and backend/docs/FRONTEND_HANDOFF.md for implemented paths and fields. Older PATCH sketches in this design are not the runtime API contract. Credential passwords are displayed only to the authorized issuer and kept out of browser storage. Public tracking tokens authorize only the associated receipt.',
281:'Validation baseline: backend compilation and 25 policy/replica-set tests passed; frontend compilation and 24 domain tests passed before the navigation update. Mobile-viewport browser tests passed representative submission, independent focal approval, credential issuance and administrator MFA/password change. The numbered design acceptance scenarios below remain a broader plan; they have not all been executed. Routing accuracy, load targets and physical old-Android acceptance are unverified.',
292:'Current execution is local: Vite on 127.0.0.1:5173, Express on loopback port 4000, a separate worker and the configured persistent MongoDB replica set. No production hosting is deployed. Select supervised production hosting, HTTPS, backup and recovery operations only when the deployment scope changes.'},
'DiscoverSalah_SRS_Phase_1.docx':{
65:'Current scope is a local application with a connected persistent MongoDB replica set. External Google Maps URLs are selected; no billed Maps/Routes API is enabled. Launch city, real mosque/staff onboarding and production operations remain owner responsibilities. Whole City uses administrative city membership; Nearby uses confirmed coordinates and a radius.',
149:'Current acceptance: changing Walking/Motorcycle changes the external Maps URL to walking/two-wheeler. The app never relabels a car or bicycle route as motorcycle. No travel duration or arrive-before-Jama’at ranking is calculated locally; timetable browsing remains available.',
216:'FR-01  Admin alone issues focal-person credentials. The responsible focal person alone issues representative credentials within assigned areas. Validate role, current scope and independently confirmed contact details on every command. Public applications do not grant management privileges.',
217:'FR-02  Authenticate with contact number and password; require temporary-password change, issuer-controlled recovery, logout, session revocation and inactive-account denial. Administrator MFA uses local TOTP or a single-use recovery code. No verified email or SMS OTP is required.',
224:'FR-09  Support permission-based device coordinates and manual location fields. Keep precise origin in memory; allow timetable browsing without it. External navigation receives only the mosque entrance, with starting point chosen in Google Maps.',
227:'FR-12  Current release opens Google Maps URLs with validated published entrance coordinates and walking/two-wheeler mode. No key, billing account or embedded map/routing SDK is required. Provider-derived route duration and distance are deferred.',
233:'FR-13  Current release does not calculate arrival margins, confident reachability or route-duration rankings. Display timetable status and explain that Google Maps handles directions separately. The earlier arrival algorithm is retained only as deferred design pending a validated routing provider.',
234:'FR-14  Display public mosque identity, timings, confidence and facilities. Provide external Google Maps directions even if no session exists for the selected date. Do not expose private Imam contacts or visitor origin in the URL.',
236:'FR-16  Future-departure forecasts remain excluded from the current local release. Do not present external navigation as an in-app forecast or promise of reaching Jama’at.'},
'Screen Inventory.docx':{
43:'Select a city and enter location details or use permission-based device coordinates. Precise location stays in memory. No geocoding API or interactive pin map is enabled; browsing remains possible without a precise origin.',
47:'Mosque/session cards show Jama’at time, date, area and timing status. No route duration or arrival margin is calculated. Preserve selected filters. Current workspace reads are bounded; complete pagination remains required before larger coverage.',
50:'Location Results and External Navigation',
51:'Keep the same result selection and venue cards, using a location list instead of simulated map geography. A selected published mosque opens the external directions screen. No embedded map SDK, tiles or clustering is loaded in the current release.',
78:'Session Details',
79:'Explain the selected congregation date/time and timing evidence. Link to directions while explicitly stating that DiscoverSalah does not calculate travel duration or an arrive-before-Jama’at estimate.',
82:'Travel Mode Selection',
83:'Choose Walking or Motorcycle for the external navigation URL. Preparation-margin calculations are deferred. Use walking or two-wheeler; do not silently substitute driving or bicycling.',
87:'Open Google Maps only after a click, using api=1, validated published entrance coordinates and walking/two-wheeler mode. Omit visitor origin, contact data and keys; use opener/referrer protection. No timetable/session is required for this screen. Revalidate app data on return.',
147:'Sign in with the issuer-confirmed contact number and password. Administrator completes MFA; temporary-password users change their password before workspace access. Public discovery remains accessible.',
150:'Issued Credentials and First Password Change',
151:'The authorized issuer privately provides a generated temporary password valid for 24 hours. Sign in, then change it before management access. No invitation token or SMS OTP activation is required; the legacy invitation page explains the issuer process.',
155:'Administrator only: verify a time-based authenticator code or unused recovery code after password sign-in. Codes are checked locally with no SMS/email provider fee.',
159:'Explain issuer-assisted reset: representative contacts the responsible focal person; focal person contacts Admin. Identity/contact are independently checked. Administrator recovery uses offline recovery material or a controlled operator process.',
163:'An authenticated owner enters the current/temporary password and a new private password. A successful change revokes prior sessions and requires sign-in again. An issuer reset creates another 24-hour temporary password.',
298:'Create Representative Credentials and Assign',
299:'The responsible focal person creates representative credentials only within their areas after independent identity/contact checks. Show the temporary password for private handoff; assign mosque responsibility through authorized commands. No SMS OTP is required.',
343:'Search scoped accounts, inspect authority, deactivate access and perform authorized resets/contact changes. Admin creates focal credentials; the responsible focal person creates representative credentials. Session revocation accompanies resets and contact changes.',
351:'Record contact/reference checks, availability, suitability, role, scope and disposition. Approval alone grants no authority. The appropriate issuer subsequently creates credentials under the admin-to-focal and focal-to-representative hierarchy.'}
}
summary=[
('Current release baseline','22 October 2026'),
('Scope','This revision governs the current local system. The existing requirement and screen identifiers are retained. Earlier diagrams, endpoint sketches and routing algorithms are historical or deferred wherever they conflict with this baseline.'),
('Navigation','External Google Maps URLs are the selected no-key approach. A user click sends only a published mosque entrance and walking or two-wheeler mode. No embedded map, Google API billing, visitor origin, private contact or route-duration response is involved. Directions work without a scheduled session. The app makes no arrival-time or reachability calculation.'),
('Accounts and trust','Only Admin issues focal credentials; only the responsible focal person issues representative credentials. Contact details are confirmed independently by the issuer. Temporary passwords expire after 24 hours and require a first-login change. No SMS OTP or verified-email recovery is required. Administrator MFA uses free local TOTP and single-use recovery codes. Imam phone-call evidence remains a separate timetable-verification process.'),
('Implementation','React/TypeScript frontend, Express API and a separate worker use persistent MongoDB. Authorization, revisions, audit and outbox writes are enforced on the server. Public responses exclude private contacts. Local ports are 5173 for the frontend and 4000 for the API; production hosting is not provisioned.'),
('Acceptance boundaries','The core local flows have automated and mobile-viewport browser coverage. Large-scale pagination, full Urdu copy review, combined registration resubmission, production recovery/load acceptance and physical old-Android validation remain outstanding. Google API routing and forecasts are deferred. See Documents/CURRENT_SYSTEM.md and backend/docs/IMPLEMENTATION_REPORT.md for the maintained implementation record.')]
manifest=[]
for name,changes in patches.items():
 path=ROOT/'Documents'/name
 original=backup/name
 if original.exists(): raise RuntimeError('Backup already exists; do not overwrite originals')
 shutil.copy2(path,original)
 doc=Document(path)
 ps=list(doc.paragraphs)+[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]
 for index,text in changes.items():
  p=ps[index];old=p.text
  p.text=text
  doc.add_comment(p.runs[0],text='22 October 2026 implementation revision. This replaces the earlier planned behavior with the current local release. The original document is preserved in .project-reference/original-documents-2026-10-22.',author='DiscoverSalah',initials='DS')
  manifest.append({'file':name,'index':index,'before':old,'after':text})
 style=doc.styles.add_style('Current Release Body',WD_STYLE_TYPE.PARAGRAPH)
 style.font.name='Calibri';style.font.size=Pt(11);style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.space_after=Pt(8);style.paragraph_format.line_spacing=1.08
 first=doc.paragraphs[0]
 for i,(label,body) in enumerate(summary):
  p=first.insert_paragraph_before(style='Current Release Body')
  if i==0:
   p.style=doc.styles['Title'];r=p.add_run(label);r.font.name='Calibri';r.font.size=Pt(22);r.font.color.rgb=RGBColor(0,0,0)
   p=first.insert_paragraph_before(body,style='Current Release Body')
  else:
   p.add_run(label+'. ').bold=True;p.add_run(body)
 first.paragraph_format.page_break_before=True
 doc.core_properties.comments='Current local release revision 22 October 2026; originals preserved.'
 doc.save(path)
 print(name, len(changes), 'targeted revisions')
(ROOT/'.project-reference/document-change-log-2026-10-22.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
